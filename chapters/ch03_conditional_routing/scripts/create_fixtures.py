"""从同一份自制 TXT 生成 DOCX、文本 PDF 与无文本 PDF 样例。"""

from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from pypdf import PdfWriter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas


def main():
    root = Path(__file__).resolve().parents[1] / "src/chapter_demo/fixtures"
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", (root / "travel_policy.txt").read_text(encoding="utf-8").strip()) if part.strip()]
    document = Document()
    document.core_properties.author = "LangGraph 学习示例"
    document.core_properties.title = paragraphs[0]
    for style_name in ("Normal", "Title"):
        style = document.styles[style_name]
        style.font.name = "宋体"
        style.font.size = Pt(12 if style_name == "Normal" else 18)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            style.element.get_or_add_rPr().append(fonts)
        fonts.set(qn("w:eastAsia"), "宋体")
    document.add_paragraph(paragraphs[0], style="Title")
    for paragraph in paragraphs[1:]:
        document.add_paragraph(paragraph)
    for element in [document.element, *(style.element for style in document.styles)]:
        for border in element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    document.save(root / "travel_policy.docx")

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    pdf = canvas.Canvas(str(root / "travel_policy.pdf"), pagesize=(595, 842), invariant=1)
    pdf.setTitle(paragraphs[0])
    pdf.setAuthor("LangGraph learning example")
    for index, paragraph in enumerate(paragraphs):
        pdf.setFont("STSong-Light", 18 if index == 0 else 12)
        pdf.drawString(55, 775 - index * 45, paragraph)
    pdf.save()
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.write(root / "no_text.pdf")
    print("已生成自制 DOCX、文本 PDF 与无文本 PDF 样例。")


if __name__ == "__main__":
    main()
