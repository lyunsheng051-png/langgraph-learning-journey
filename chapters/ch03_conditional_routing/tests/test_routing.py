"""用真实图执行验证条件分支、来源结构与失败时的停止行为。"""

from contextlib import ExitStack
from importlib.resources import as_file, files
from pathlib import Path
from tempfile import gettempdir, TemporaryDirectory
import unittest
from unittest.mock import patch

from docx import Document
from pypdf import PdfReader, PdfWriter

from chapter_demo import graph as graph_module


class ConditionalRoutingTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory(prefix="langgraph-ch03-")
        self.root = Path(self.directory.name).resolve()
        self.assertEqual(self.root.parent, Path(gettempdir()).resolve())
        self.addCleanup(self.directory.cleanup)

    def run_file(self, path):
        return graph_module.build_graph().invoke({"input_path": str(path), "logs": []})

    def fixture_path(self, stack, name):
        return stack.enter_context(as_file(files("chapter_demo").joinpath("fixtures", name)))

    def assert_rejected_without_analysis(self, path, reason):
        with patch.object(graph_module, "analyze_document", side_effect=AssertionError("失败输入不应统计")):
            result = self.run_file(path)
        self.assertEqual(result["status"], "error")
        self.assertIn(reason, result["error"])
        self.assertNotIn("character_count", result)
        return result

    def test_supported_types_select_only_the_matching_extractor(self):
        node_names = {"txt": "extract_txt_node", "docx": "extract_docx_node", "pdf": "extract_pdf_node"}
        for kind, selected in node_names.items():
            with self.subTest(kind=kind), ExitStack() as stack:
                for name in node_names.values():
                    if name != selected:
                        stack.enter_context(patch.object(graph_module, name, side_effect=AssertionError("不应运行其他格式分支")))
                path = self.fixture_path(stack, f"travel_policy.{kind}")
                result = self.run_file(path)
                self.assertEqual(result["status"], "ok")
                self.assertEqual(result["character_count"], 60)
                self.assertEqual(result["file_type"], kind)
                self.assertEqual(len(result["logs"]), 4)
                self.assertEqual(result["section_count"], 1 if kind == "pdf" else 3)

    def test_unsupported_type_never_opens_an_extractor(self):
        path = self.root / "example.md"
        path.write_text("text", encoding="utf-8")
        with patch.object(graph_module, "extract_with", side_effect=AssertionError("不支持格式不应提取")):
            result = self.assert_rejected_without_analysis(path, "不支持文件类型")
        self.assertEqual(len(result["logs"]), 2)

    def test_pdf_without_text_is_rejected_after_extraction(self):
        with ExitStack() as stack:
            path = self.fixture_path(stack, "no_text.pdf")
            result = self.assert_rejected_without_analysis(path, "未提取到可处理文本")
        self.assertIn("提取失败（PDF）", result["logs"][1])

    def test_pdf_keeps_original_page_numbers_when_blank_pages_are_skipped(self):
        with ExitStack() as stack:
            source = self.fixture_path(stack, "travel_policy.pdf")
            writer = PdfWriter()
            writer.add_blank_page(width=595, height=842)
            writer.add_page(PdfReader(source).pages[0])
            path = self.root / "second-page.pdf"
            writer.write(path)
        result = self.run_file(path)
        self.assertEqual(result["sections"][0]["location"], "第 2 页")
        self.assertEqual(result["section_count"], 1)

    def test_docx_keeps_original_paragraph_numbers(self):
        document = Document()
        document.add_paragraph("")
        document.add_paragraph("甲乙。")
        path = self.root / "paragraphs.docx"
        document.save(path)
        result = self.run_file(path)
        self.assertEqual(result["sections"], [{"location": "第 2 段", "text": "甲乙。"}])

    def test_missing_file_does_not_enter_analysis(self):
        self.assert_rejected_without_analysis(self.root / "missing.txt", "文件不存在")

    def test_empty_txt_and_docx_are_rejected(self):
        text = self.root / "empty.txt"
        text.write_text(" \n\t", encoding="utf-8")
        self.assert_rejected_without_analysis(text, "未提取到可处理文本")
        word = self.root / "empty.docx"
        Document().save(word)
        self.assert_rejected_without_analysis(word, "未提取到可处理文本")

    def test_malformed_and_renamed_files_are_not_treated_as_converted(self):
        for suffix in ("docx", "pdf"):
            with self.subTest(suffix=suffix):
                path = self.root / f"renamed.{suffix}"
                path.write_bytes(b"this is plain text, not a document package")
                self.assert_rejected_without_analysis(path, "无法解析")

    def test_invalid_utf8_is_a_routed_error(self):
        path = self.root / "invalid.txt"
        path.write_bytes(b"\xff")
        self.assert_rejected_without_analysis(path, "不是有效的 UTF-8")

    def test_uppercase_extension_and_bom_work(self):
        path = self.root / "中文 空格.TXT"
        path.write_bytes("\ufeff甲乙。\r\n\r\n丙。".encode("utf-8"))
        result = self.run_file(path)
        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["character_count"], 5)
        self.assertEqual(result["section_count"], 2)

    def test_encrypted_pdf_and_page_limit_are_reported(self):
        encrypted = self.root / "encrypted.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        writer.encrypt("example-password")
        writer.write(encrypted)
        self.assert_rejected_without_analysis(encrypted, "PDF 已加密")
        many_pages = self.root / "many-pages.pdf"
        writer = PdfWriter()
        for _ in range(21):
            writer.add_blank_page(width=595, height=842)
        writer.write(many_pages)
        self.assert_rejected_without_analysis(many_pages, "20 页上限")

    def test_size_and_text_limits_stop_before_analysis(self):
        oversized = self.root / "oversized.txt"
        oversized.write_bytes(b"x" * (1024 * 1024 + 1))
        self.assert_rejected_without_analysis(oversized, "1 MiB 上限")
        too_long = self.root / "long.txt"
        too_long.write_bytes(b"x" * 100_001)
        self.assert_rejected_without_analysis(too_long, "100,000 字符上限")

    def test_reusing_graph_for_new_input_does_not_keep_an_old_error(self):
        graph = graph_module.build_graph()
        failed = graph.invoke({"input_path": str(self.root / "missing.txt"), "logs": []})
        path = self.root / "new.txt"
        path.write_text("新资料", encoding="utf-8")
        succeeded = graph.invoke({"input_path": str(path), "logs": []})
        self.assertEqual(failed["status"], "error")
        self.assertEqual(succeeded["status"], "ok")
        self.assertNotIn("error", succeeded)
        self.assertEqual(len(succeeded["logs"]), 4)


if __name__ == "__main__":
    unittest.main()
